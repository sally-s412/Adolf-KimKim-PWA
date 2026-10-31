import importlib.util
from io import BytesIO
from pathlib import Path
import unittest

from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject

spec = importlib.util.spec_from_file_location('document_routes', Path(__file__).resolve().parents[1] / 'routes/documents.py')
documents = importlib.util.module_from_spec(spec)
spec.loader.exec_module(documents)


def pdf_bytes(text=None, password=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 12 Tf 20 300 Td (' + text + ') Tj ET').encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class DocumentUploads(unittest.TestCase):
    def setUp(self):
        app = FastAPI()
        app.include_router(documents.router)
        self.client = TestClient(app)

    def upload(self, content, name, language='vi'):
        return self.client.post('/api/documents/extract', files={'file': (name, content)}, headers={'Accept-Language': language})

    def test_docx_keeps_vietnamese_and_table_in_document_order(self):
        document = Document()
        document.add_paragraph('Nội dung mở đầu')
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Doanh thu'
        document.tables[0].cell(0, 1).text = '200 triệu'
        document.add_paragraph('Kết luận cuối')
        output = BytesIO()
        document.save(output)
        response = self.upload(output.getvalue(), 'báo cáo.docx')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual([row['text'] for row in data['sections']], ['Nội dung mở đầu', 'Doanh thu | 200 triệu', 'Kết luận cuối'])
        self.assertIsNone(data['page_count'])

    def test_pdf_extracts_actual_text_and_page_number(self):
        response = self.upload(pdf_bytes('Original document text'), 'report.pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['sections'][0]['text'], 'Original document text')
        self.assertEqual(response.json()['sections'][0]['page'], 1)

    def test_image_only_pdf_reports_need_for_ocr(self):
        response = self.upload(pdf_bytes(), 'scan.pdf', 'en')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()['detail']['code'], 'empty')
        self.assertIn('OCR', response.json()['detail']['message'])

    def test_password_protected_pdf_is_not_silently_imported(self):
        response = self.upload(pdf_bytes('Private text', password='test-password'), 'locked.pdf')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()['detail']['code'], 'locked')

    def test_legacy_doc_and_fake_docx_are_rejected_clearly(self):
        self.assertEqual(self.upload(b'legacy', 'report.doc').status_code, 415)
        response = self.upload(b'not a Word file', 'report.docx')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()['detail']['code'], 'invalid')

    def test_oversize_document_fails_before_parsing(self):
        response = self.upload(b'x' * (documents.MAX_BYTES + 1), 'large.pdf')
        self.assertEqual(response.status_code, 413)


if __name__ == '__main__':
    unittest.main()
